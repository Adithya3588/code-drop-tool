from __future__ import annotations

import io
import json
import subprocess
import tempfile
import time
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from app import create_app
from release_uploader import (
    WorkflowError,
    compare_binaries,
    compare_release_history,
    config_diffs,
    extract_configs_zip,
    extract_zip,
    load_json,
    scan_configs,
    run_svn,
    validate_configs,
    validate_folder_name,
    validate_svn_url,
)
from reporting import build_binaries_report, build_configs_report


class WorkflowHelpersTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def make_zip(self, files: dict[str, bytes]) -> io.BytesIO:
        result = io.BytesIO()
        with zipfile.ZipFile(result, "w") as archive:
            for name, content in files.items():
                archive.writestr(name, content)
        result.seek(0)
        return result

    def test_input_validation(self):
        self.assertEqual(validate_folder_name("Codedrop10"), "Codedrop10")
        self.assertEqual(validate_folder_name("Codedrop8"), "Codedrop8")
        with self.assertRaises(WorkflowError):
            validate_folder_name("release-final")
        self.assertEqual(validate_svn_url("https://svn.example/repo/CodeDrops"), "https://svn.example/repo/CodeDrops")
        with self.assertRaises(WorkflowError):
            validate_svn_url("https://user:password@svn.example/repo")

    def test_svn_password_is_sent_via_stdin_and_not_cached(self):
        completed = subprocess.CompletedProcess([], 0, "ok", "")
        with patch("release_uploader.svn_executable", return_value="svn.exe"), patch(
            "release_uploader.subprocess.run", return_value=completed
        ) as invoke:
            run_svn(["info", "https://svn.example/repo"], "release-user", "temporary-secret")
        arguments = invoke.call_args.args[0]
        self.assertIn("--no-auth-cache", arguments)
        self.assertIn("--password-from-stdin", arguments)
        self.assertNotIn("temporary-secret", arguments)
        self.assertEqual(invoke.call_args.kwargs["input"], "temporary-secret\n")

    def test_zip_extracts_wrapped_code_drop_and_blocks_traversal(self):
        archive = self.make_zip({
            "Codedrop10/Configs/Api/config.json": b'{"enabled": true}',
            "Codedrop10/Binaries/Api/service.dll": b"binary",
        })
        root = extract_zip(archive, self.root / "current")
        self.assertEqual(root.name, "Codedrop10")
        self.assertTrue((root / "Configs/Api/config.json").is_file())
        malicious = self.make_zip({"../outside.txt": b"bad"})
        with self.assertRaises(WorkflowError):
            extract_zip(malicious, self.root / "unsafe")
        self.assertFalse((self.root / "outside.txt").exists())

    def test_configs_zip_accepts_contents_without_configs_wrapper(self):
        archive = self.make_zip({"Api/settings.json": b'{"port": 8080}'})
        extracted = extract_configs_zip(archive, self.root / "config-upload")
        self.assertEqual(json.loads((extracted / "Api/settings.json").read_text()), {"port": 8080})

    def test_json_validation_reports_every_invalid_file_and_duplicate_warning(self):
        configs = self.root / "Configs"
        configs.mkdir()
        (configs / "bad.json").write_text('{\n  "broken":\n}', encoding="utf-8")
        (configs / "duplicates.json").write_text('{"api": {"port": 1, "port": 2}}', encoding="utf-8")
        report = validate_configs(configs)
        self.assertEqual(report["checked"], 2)
        self.assertEqual(report["failed"], 1)
        self.assertEqual(report["errors"][0]["file"], "bad.json")
        self.assertGreater(report["errors"][0]["line"], 0)
        self.assertEqual(report["duplicates"][0]["key_path"], "api.port")
        parsed, duplicates = load_json(configs / "duplicates.json")
        self.assertEqual(parsed["api"]["port"], 2)
        self.assertEqual(len(duplicates), 1)

    def test_scan_configs_skips_invalid_json_files_instead_of_crashing(self):
        configs = self.root / "Configs"
        configs.mkdir()
        (configs / "valid.json").write_text('{"password": "secret"}', encoding="utf-8")
        (configs / "broken.json").write_text('{"bad": }', encoding="utf-8")
        rules = self.root / "rules.json"
        rules.write_text(json.dumps({
            "version": 1,
            "ignoreCase": True,
            "rules": [
                {"pattern": "secret", "match": "contains", "severity": "Warning"},
            ],
            "emptyRequiredKeys": [],
            "allowlist": [],
        }), encoding="utf-8")
        findings, rules_hash = scan_configs(configs, rules)
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["file"], "valid.json")
        self.assertEqual(findings[0]["value"], "••••••••")
        self.assertEqual(len(rules_hash), 64)

    def test_config_diff_includes_json_parse_context_for_invalid_files(self):
        current_root = self.root / "Codedrop10"
        current_cfg = current_root / "Configs" / "Api"
        current_cfg.mkdir(parents=True)
        (current_cfg / "settings.json").write_text('{"bad": }', encoding="utf-8")

        baseline_root = self.root / "Codedrop9" / "Configs" / "Api"
        baseline_root.mkdir(parents=True)
        (baseline_root / "settings.json").write_text('{"ok": true}', encoding="utf-8")

        results = config_diffs(current_root, {
            "Configs/Api": {
                "code_drop": "Codedrop9",
                "local_path": str(baseline_root),
                "path": "Configs/Api",
            }
        })

        self.assertEqual(results[0]["status"], "Invalid JSON")
        self.assertIn("context", results[0]["current_validation_error"])
        self.assertIn("bad", results[0]["current_validation_error"]["context"])

    def test_dummy_scan_masks_secret_values_and_applies_configured_rules(self):
        configs = self.root / "Configs"
        configs.mkdir()
        (configs / "Api/settings.json").parent.mkdir()
        (configs / "Api/settings.json").write_text(
            json.dumps({"Database": {"Password": "password123"}, "Name": "test service"}),
            encoding="utf-8",
        )
        rules = self.root / "rules.json"
        rules.write_text(json.dumps({
            "version": 1,
            "ignoreCase": True,
            "rules": [
                {"pattern": "password", "match": "contains", "severity": "Blocker"},
                {"pattern": "test", "match": "wholeWord", "severity": "Warning"},
            ],
            "emptyRequiredKeys": ["password", "token"],
            "allowlist": [],
        }), encoding="utf-8")
        findings, rules_hash = scan_configs(configs, rules)
        password = next(item for item in findings if item["key_path"] == "Database.Password")
        self.assertEqual(password["severity"], "Blocker")
        self.assertEqual(password["value"], "••••••••")
        self.assertEqual(len(rules_hash), 64)
        self.assertTrue(any(item["key_path"] == "Name" for item in findings))

    def test_default_dummy_rules_scan_values_not_password_key_names(self):
        configs = self.root / "Configs"
        configs.mkdir()
        (configs / "settings.json").write_text(
            json.dumps({
                "tmc.conf.POMPassword": "",
                "tmc.conf.POMUsername": "",
                "Credential": "password123",
            }),
            encoding="utf-8",
        )

        findings, _ = scan_configs(
            configs, Path(__file__).parents[1] / "config" / "dummy_rules.json"
        )

        self.assertFalse(any(item["key_path"] == "tmc.conf.POMPassword" for item in findings))
        self.assertTrue(any(
            item["key_path"] == "Credential" and item["pattern"] == "password"
            for item in findings
        ))

    def test_dummy_scan_skips_template_references_but_scans_tmc_data_values(self):
        configs = self.root / "Configs" / "TMAC_TMS_ExternalConnector" / "config" / "app"
        configs.mkdir(parents=True)
        (configs / "template.json").write_text(
            '{"HostSettings":{"CertificatePassword":"tmc.conf.CertificatePassword"},'
            '"AppSettings":{"POM":{"Password":"tmc.conf.POMPassword"}}}',
            encoding="utf-8",
        )
        (configs / "TMC_Data.json").write_text(
            '{"tmc.conf.CertificatePassword":"",'
            '"tmc.conf.POMPassword":"password123"}',
            encoding="utf-8",
        )

        findings, _ = scan_configs(
            self.root / "Configs",
            Path(__file__).parents[1] / "config" / "dummy_rules.json",
        )

        self.assertFalse(any(item["file"].endswith("template.json") for item in findings))
        self.assertEqual(
            [(item["file"], item["key_path"], item["pattern"]) for item in findings],
            [(
                "TMAC_TMS_ExternalConnector/config/app/TMC_Data.json",
                "tmc.conf.POMPassword",
                "password",
            )],
        )

    def test_dummy_scan_exercises_match_modes_component_allowlist_and_duplicates(self):
        configs = self.root / "Configs"
        component = configs / "Api"
        component.mkdir(parents=True)
        (component / "settings.json").write_text(
            '{"Mode":"TESTSERVICE","Code":"test service","Exact":"release",'
            '"Endpoint":"","Password":"changeme","Allowed":"known",'
            '"Items":[{"Name":"dummy"}],"Duplicate":1,"Duplicate":2}',
            encoding="utf-8",
        )
        rules = self.root / "rules.json"
        rules.write_text(json.dumps({
            "version": 1,
            "ignoreCase": True,
            "scanKeys": True,
            "rules": [
                {"pattern": "test", "match": "wholeWord", "severity": "Warning"},
                {"pattern": "service", "match": "contains", "severity": "Blocker"},
                {"pattern": "release", "match": "exact", "severity": "Warning"},
                {"pattern": "^dummy$", "match": "regex", "severity": "Blocker"},
                {"pattern": "changeme", "match": "contains", "severity": "Blocker", "component": "Api"},
                {"pattern": "disabled", "match": "contains", "severity": "Blocker", "enabled": False},
                {"pattern": "known", "match": "exact", "severity": "Warning"},
            ],
            "emptyRequiredKeys": ["endpoint"],
            "allowlist": [{
                "file": "Api/settings.json",
                "keyPath": "Allowed",
                "pattern": "known",
                "reason": "Approved test value",
            }],
        }, indent=2), encoding="utf-8")

        findings, rules_hash = scan_configs(configs, rules)
        by_key_pattern = {(item["key_path"], item["pattern"]): item for item in findings}
        self.assertIn(("Code", "test"), by_key_pattern)
        self.assertIn(("Mode", "service"), by_key_pattern)
        self.assertIn(("Exact", "release"), by_key_pattern)
        self.assertIn(("Items[0].Name", "^dummy$"), by_key_pattern)
        self.assertIn(("Password", "changeme"), by_key_pattern)
        self.assertIn(("Endpoint", "empty required value"), by_key_pattern)
        self.assertIn(("Duplicate", "duplicate JSON key"), by_key_pattern)
        self.assertNotIn(("Allowed", "known"), by_key_pattern)
        self.assertFalse(any(item["pattern"] == "disabled" for item in findings))
        self.assertEqual(by_key_pattern[("Password", "changeme")]["value"], "••••••••")
        self.assertEqual(len(rules_hash), 64)

    def test_config_semantic_diff_includes_additions_and_removed_files(self):
        current = self.root / "current"
        configs = current / "Configs" / "Api"
        configs.mkdir(parents=True)
        (configs / "settings.json").write_text(
            '{"Database":{"Password":"safe"},"Port":9000}', encoding="utf-8"
        )
        baseline = self.root / "baseline"
        baseline.mkdir()
        (baseline / "settings.json").write_text(
            '{"Database":{"Password":"old-secret"},"Port":8000,"Enabled":true}', encoding="utf-8"
        )
        (baseline / "removed.json").write_text('{"Old":true}', encoding="utf-8")
        results = config_diffs(current, {
            "Configs/Api": {
                "code_drop": "Codedrop09",
                "path": "Configs/Api",
                "revision": "17",
                "local_path": str(baseline),
            }
        })
        settings = next(item for item in results if item["file"] == "Api/settings.json")
        changes = {item["key_path"]: item for item in settings["changes"]}
        self.assertEqual(changes["Port"]["change"], "Changed")
        self.assertEqual(changes["Database.Password"]["previous"], "••••••••")
        self.assertEqual(changes["Enabled"]["change"], "Removed")
        removed = next(item for item in results if item["file"] == "Api/removed.json")
        self.assertEqual(removed["status"], "Removed")

    def test_invalid_baseline_json_is_reported_without_exposing_file_contents(self):
        current = self.root / "release" / "Configs" / "Api"
        baseline = self.root / "baseline" / "Api"
        current.mkdir(parents=True)
        baseline.mkdir(parents=True)
        (current / "settings.json").write_text('{"token":"current-safe"}', encoding="utf-8")
        (baseline / "settings.json").write_text(
            '{"token":"baseline-secret", invalid}', encoding="utf-8"
        )
        (baseline / "removed.json").write_text(
            '{"password":"do-not-print", invalid}', encoding="utf-8"
        )

        results = config_diffs(self.root / "release", {
            "Configs/Api": {
                "code_drop": "Codedrop09",
                "path": "Configs/Api",
                "revision": "9",
                "local_path": str(baseline),
            }
        })
        by_file = {item["file"]: item for item in results}
        self.assertEqual(by_file["Api/settings.json"]["status"], "Baseline invalid JSON")
        self.assertGreater(by_file["Api/settings.json"]["baseline_validation_error"]["line"], 0)
        self.assertEqual(by_file["Api/removed.json"]["status"], "Removed")
        self.assertIn("baseline_validation_error", by_file["Api/removed.json"])
        self.assertNotIn("baseline-secret", str(by_file))
        self.assertNotIn("do-not-print", str(by_file))
        history = compare_release_history(self.root / "release", {
            "Configs/Api": [
                {
                    "code_drop": "Codedrop09",
                    "local_path": str(baseline),
                },
            ]
        })
        self.assertTrue(all(row["json_key_changes"] != "Invalid JSON" for row in history))

    def test_config_diff_covers_all_file_outcomes_and_redacts_sensitive_json(self):
        current = self.root / "release" / "Configs" / "Api"
        baseline = self.root / "baseline" / "Api"
        current.mkdir(parents=True)
        baseline.mkdir(parents=True)
        (baseline / "same.json").write_text('{"enabled": true}', encoding="utf-8")
        (current / "same.json").write_text('{"enabled": true}', encoding="utf-8")
        (baseline / "changed.json").write_text(
            '{"token":"old-secret","ports":[8000,8001],"removed":true}', encoding="utf-8"
        )
        (current / "changed.json").write_text(
            '{"token":"new-secret","ports":[8000,9000],"added":true}', encoding="utf-8"
        )
        (current / "new.json").write_text('{"new": true}', encoding="utf-8")
        (baseline / "removed.json").write_text('{"gone": true}', encoding="utf-8")
        (baseline / "settings.yaml").write_text("port: 8000\n", encoding="utf-8")
        (current / "settings.yaml").write_text("port: 9000\n", encoding="utf-8")
        (current / "new.properties").write_text("mode=production\n", encoding="utf-8")
        (baseline / "removed.bin").write_bytes(b"\x00old")

        results = config_diffs(self.root / "release", {
            "Configs/Api": {
                "code_drop": "Codedrop09",
                "path": "Configs/Api",
                "revision": "9",
                "local_path": str(baseline),
            }
        })
        by_file = {item["file"]: item for item in results}
        self.assertEqual(by_file["Api/same.json"]["status"], "Unchanged")
        self.assertEqual(by_file["Api/changed.json"]["status"], "Changed")
        changes = {item["key_path"]: item for item in by_file["Api/changed.json"]["changes"]}
        self.assertEqual(changes["ports[1]"]["change"], "Changed")
        self.assertEqual(changes["added"]["change"], "Added")
        self.assertEqual(changes["removed"]["change"], "Removed")
        self.assertEqual(changes["token"]["previous"], "••••••••")
        self.assertEqual(changes["token"]["current"], "••••••••")
        self.assertNotIn("old-secret", str(by_file["Api/changed.json"]["text_diff"]))
        self.assertEqual(by_file["Api/new.json"]["status"], "New")
        self.assertEqual(by_file["Api/removed.json"]["status"], "Removed")
        self.assertEqual(by_file["Api/settings.yaml"]["status"], "Changed")
        self.assertEqual(by_file["Api/new.properties"]["status"], "New")
        self.assertEqual(by_file["Api/removed.bin"]["status"], "Removed")
        for file in ("Api/settings.yaml", "Api/new.properties", "Api/removed.bin"):
            self.assertRegex(by_file[file]["changes"][0]["current"], r"^[0-9a-f]{64}$|^–$")

    def test_binary_comparison_detects_changes_and_new_files(self):
        current = self.root / "current" / "Binaries" / "Api"
        baseline = self.root / "baseline"
        current.mkdir(parents=True)
        baseline.mkdir()
        (current / "service.dll").write_bytes(b"new binary")
        (baseline / "service.dll").write_bytes(b"old binary")
        (current / "new.exe").write_bytes(b"new file")
        rows = compare_binaries(self.root / "current", {
            "Binaries/Api": {
                "code_drop": "Codedrop09",
                "path": "Binaries/Api",
                "revision": "17",
                "local_path": str(baseline),
            }
        })
        self.assertEqual({row["status"] for row in rows}, {"Changed", "New"})
        self.assertIn("N/A", {row["baseline_version"] for row in rows})

    def test_binary_comparison_covers_unchanged_rebuilt_moved_removed_and_version_changes(self):
        current_component = self.root / "release" / "Binaries" / "Api"
        baseline_component = self.root / "baseline" / "Api"
        current_component.mkdir(parents=True)
        baseline_component.mkdir(parents=True)
        contents = {
            "unchanged.exe": (b"same", b"same"),
            "rebuild.exe": (b"old-build", b"new-build"),
            "upgrade.exe": (b"upgrade-old", b"upgrade-new"),
            "downgrade.exe": (b"downgrade-old", b"downgrade-new"),
        }
        for name, (old, new) in contents.items():
            (baseline_component / name).write_bytes(old)
            (current_component / name).write_bytes(new)
        (baseline_component / "old-name.exe").write_bytes(b"moved-binary")
        (current_component / "nested").mkdir()
        (current_component / "nested" / "new-name.exe").write_bytes(b"moved-binary")
        (baseline_component / "removed.exe").write_bytes(b"removed")
        (current_component / "new.exe").write_bytes(b"new")

        def version(path: Path) -> str:
            versions = {
                "unchanged.exe": "1.0.0.0",
                "rebuild.exe": "1.0.0.0",
                "upgrade.exe": "2.0.0.0" if path.parent == current_component else "1.0.0.0",
                "downgrade.exe": "1.0.0.0" if path.parent == current_component else "2.0.0.0",
                "old-name.exe": "3.0.0.0",
                "new-name.exe": "3.0.0.0",
                "removed.exe": "1.0.0.0",
                "new.exe": "1.0.0.0",
            }
            return versions[path.name]

        with patch("release_uploader._binary_version", side_effect=version):
            rows = compare_binaries(self.root / "release", {
                "Binaries/Api": {
                    "code_drop": "Codedrop09",
                    "path": "Binaries/Api",
                    "revision": "9",
                    "local_path": str(baseline_component),
                }
            })
        status_by_file = {row["file"]: row["status"] for row in rows}
        self.assertEqual(status_by_file, {
            "downgrade.exe": "Version downgrade",
            "nested/new-name.exe": "Moved",
            "new.exe": "New",
            "rebuild.exe": "Hash changed, same version",
            "removed.exe": "Removed",
            "unchanged.exe": "Unchanged",
            "upgrade.exe": "Changed",
        })

    def test_historical_comparison_reports_missing_and_changed_files_for_each_release(self):
        current = self.root / "current"
        current_component = current / "Binaries" / "Api"
        current_component.mkdir(parents=True)
        (current_component / "service.exe").write_bytes(b"current")
        baseline9 = self.root / "Codedrop9" / "Binaries" / "Api"
        baseline8 = self.root / "Codedrop8" / "Binaries" / "Api"
        baseline9.mkdir(parents=True)
        baseline8.mkdir(parents=True)
        (baseline9 / "service.exe").write_bytes(b"release9")
        (baseline9 / "retired.dll").write_bytes(b"retired")
        (baseline8 / "service.exe").write_bytes(b"current")
        rows = compare_release_history(current, {
            "Binaries/Api": [
                {"code_drop": "Codedrop9", "local_path": str(baseline9)},
                {"code_drop": "Codedrop8", "local_path": str(baseline8)},
            ],
        })
        outcomes = {
            (row["code_drop"], row["file"]): row["status"]
            for row in rows
        }
        self.assertEqual(outcomes[("Codedrop9", "service.exe")], "Changed")
        self.assertEqual(outcomes[("Codedrop9", "retired.dll")], "Missing from current")
        self.assertEqual(outcomes[("Codedrop8", "service.exe")], "Unchanged")

    def test_workbooks_include_expected_sheets_and_diff_image(self):
        binary_path = self.root / "Binaries.xlsx"
        config_path = self.root / "Configs.xlsx"
        build_binaries_report(binary_path, "Codedrop10", [{
            "component": "Api",
            "file": '=HYPERLINK("https://invalid")',
            "baseline_path": "No baseline found",
            "baseline_version": "–",
            "current_version": "N/A",
            "status": "New",
            "baseline_drop": "",
        }], history_releases=["Codedrop09", "Codedrop10"], version_history=[{
            "component": "Api",
            "file": "service.exe",
            "baseline_drop": "Codedrop09",
            "versions": {"Codedrop09": "1.0.0.0", "Codedrop10": "2.0.0.0"},
        }])
        build_configs_report(config_path, "Codedrop10", [{
            "file": "Api/settings.json",
            "status": "Changed",
            "baseline": {"code_drop": "Codedrop09", "path": "Configs/Api"},
            "changes": [{
                "key_path": "Port",
                "previous": "8000",
                "current": "9000",
                "change": "Changed",
            }],
        }], [], "a" * 64, {
            "checked": 1, "valid": 1, "failed": 0, "errors": [], "duplicates": [],
        }, True, "completed")
        from openpyxl import load_workbook

        binaries = load_workbook(binary_path, read_only=True)
        self.assertEqual(binaries.sheetnames, [
            "Binary Comparison", "Version History",
            "All Release Comparisons", "Release Coverage",
        ])
        self.assertNotIn("Baselines", binaries.sheetnames)
        binary_file_cell = binaries["Binary Comparison"]["B2"]
        self.assertEqual(binary_file_cell.data_type, "s")
        self.assertTrue(binary_file_cell.value.startswith("'="))
        history_sheet = binaries["Version History"]
        self.assertEqual(history_sheet["C1"].value, "Codedrop09")
        self.assertEqual(history_sheet["D1"].value, "Codedrop10 (current)")
        self.assertEqual(history_sheet["C2"].value, "1.0.0.0")
        self.assertEqual(history_sheet["D2"].value, "2.0.0.0")
        binaries.close()
        configs = load_workbook(config_path)
        self.assertIn("Scan Summary", configs.sheetnames)
        self.assertIn("Validation Issues", configs.sheetnames)
        scan_summary = {
            row[0].value: row[1].value
            for row in configs["Scan Summary"].iter_rows(min_row=2)
        }
        self.assertEqual(scan_summary["Dummy scan status"], "completed")
        self.assertIn("Api_settings", configs.sheetnames)
        self.assertEqual(len(configs["Api_settings"]._images), 1)
        sheet = configs["Api_settings"]
        values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
        self.assertIn("Key path", values)
        self.assertIn("Port", values)
        configs.close()


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.app = create_app({
            "TESTING": True,
            "DATABASE": str(root / "db.sqlite3"),
            "WORKSPACE": str(root / "work"),
            "RULES_FILE": str(Path(__file__).parents[1] / "config" / "dummy_rules.json"),
        })
        self.client = self.app.test_client()

    def tearDown(self):
        self.app.extensions["release_review_executor"].shutdown(wait=True)
        self.temp.cleanup()

    def test_health_and_session_input_validation(self):
        health = self.client.get("/api/health")
        self.assertEqual(health.status_code, 200)
        self.assertEqual(health.get_json()["database"], "SQLite")
        invalid = self.client.post("/api/sessions", json={
            "folder_name": "release-10",
            "target_url": "https://svn.example/repo/CodeDrops",
        })
        self.assertEqual(invalid.status_code, 400)
        valid = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": "https://svn.example/repo/CodeDrops",
        })
        self.assertEqual(valid.status_code, 201)
        self.assertTrue(valid.get_json()["id"])

    def test_reports_are_not_available_before_a_source_is_analyzed(self):
        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": "https://svn.example/repo/CodeDrops",
        }).get_json()
        response = self.client.get(f"/api/sessions/{created['id']}/reports/Configs.xlsx")
        self.assertEqual(response.status_code, 404)

    def test_remote_svn_source_requires_credentials(self):
        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": "https://svn.example/repo/CodeDrops",
        }).get_json()
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as zip_file:
            zip_file.writestr("Codedrop10/Binaries/Api/service.exe", b"release")
        archive.seek(0)
        response = self.client.post(
            f"/api/sessions/{created['id']}/source",
            data={"archive": (archive, "Codedrop10.zip")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("SVN username and password", response.get_json()["error"])

    def test_frontend_and_session_api_routes_work_end_to_end(self):
        page = self.client.get("/")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"Review a code drop", page.data)
        self.assertIn(b"app.css", page.data)
        self.assertIn(b"SVN URL or local folder path", page.data)
        self.assertIn(b"existing absolute local folder path", page.data)
        self.assertIn(b"SVN username", page.data)
        frontend_script = self.client.get("/static/app.js")
        self.assertEqual(frontend_script.status_code, 200)
        self.assertIn(b'["New", "Changed", "Invalid JSON", "Baseline invalid JSON"]', frontend_script.data)
        self.assertIn(b'.filter(row => row.status !== "Removed")', frontend_script.data)
        self.assertIn(b"files that exist only in earlier releases are omitted", frontend_script.data)
        self.assertNotIn(b"Rules SHA-256", frontend_script.data)
        self.assertNotIn(b"Rules are maintained", frontend_script.data)
        self.assertNotIn(b"JSON checked:", frontend_script.data)
        frontend_script.close()

        working_copy = Path(self.temp.name) / "SVN checkout"
        (working_copy / ".svn").mkdir(parents=True)
        for drop, content in (("Codedrop9", b"nine"), ("Codedrop8", b"eight")):
            binary = working_copy / drop / "Binaries" / "Api" / "service.exe"
            binary.parent.mkdir(parents=True)
            binary.write_bytes(content)

        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": str(working_copy),
        })
        self.assertEqual(created.status_code, 201)
        session_id = created.get_json()["id"]

        missing_source = self.client.post(
            f"/api/sessions/{session_id}/source",
            data={"source_mode": "zip"},
            content_type="multipart/form-data",
        )
        self.assertEqual(missing_source.status_code, 400)
        self.assertIn("Choose a ZIP", missing_source.get_json()["error"])

        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as zip_file:
            zip_file.writestr("Codedrop10/Binaries/Api/service.exe", b"ten")
            zip_file.writestr(
                "Codedrop10/Configs/Api/settings.json",
                '{"Endpoint":"abc"}',
            )
        archive.seek(0)
        started = self.client.post(
            f"/api/sessions/{session_id}/source",
            data={"archive": (archive, "Codedrop10.zip")},
            content_type="multipart/form-data",
        )
        self.assertEqual(started.status_code, 202)
        for _ in range(100):
            details = self.client.get(f"/api/sessions/{session_id}")
            if details.get_json()["status"] != "analyzing":
                break
            time.sleep(0.02)
        self.assertEqual(details.status_code, 200)
        self.assertEqual(details.get_json()["status"], "ready")
        self.assertEqual(details.get_json()["state"]["findings"][0]["severity"], "Warning")

        invalid_baseline = self.client.post(
            f"/api/sessions/{session_id}/baseline",
            json={"path": "Binaries/Api", "code_drop": "Codedrop7"},
        )
        self.assertEqual(invalid_baseline.status_code, 400)
        changed_baseline = self.client.post(
            f"/api/sessions/{session_id}/baseline",
            json={"path": "Binaries/Api", "code_drop": "Codedrop8"},
        )
        self.assertEqual(changed_baseline.status_code, 200)
        self.assertEqual(
            changed_baseline.get_json()["baselines"]["Binaries/Api"]["code_drop"],
            "Codedrop8",
        )

        for report_name in ("Binaries.xlsx", "Configs.xlsx"):
            report = self.client.get(
                f"/api/sessions/{session_id}/reports/{report_name}"
            )
            self.assertEqual(report.status_code, 200)
            report.close()
        report_bundle = self.client.get(
            f"/api/sessions/{session_id}/reports/Release-Review.zip"
        )
        self.assertEqual(report_bundle.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(report_bundle.data)) as bundle:
            self.assertEqual(
                set(bundle.namelist()),
                {
                    "Binaries.xlsx",
                    "Configs.xlsx",
                    "Codedrop10-Review-Notes.txt",
                },
            )
        report_bundle.close()
        self.assertEqual(
            self.client.get(f"/api/sessions/{session_id}/reports/unknown.xlsx").status_code,
            404,
        )
        csv_report = self.client.get(
            f"/api/sessions/{session_id}/json-errors.csv"
        )
        self.assertEqual(csv_report.status_code, 200)
        self.assertIn(b"File,Line,Column,Error,Context", csv_report.data)

        reviewed = self.client.post(
            f"/api/sessions/{session_id}/review", json={"reviewed": True}
        )
        self.assertEqual(reviewed.status_code, 200)
        finding_id = reviewed.get_json()["findings"][0]["id"]
        acknowledged = self.client.post(
            f"/api/sessions/{session_id}/acknowledge",
            json={"finding_id": finding_id, "acknowledged": True},
        )
        self.assertEqual(acknowledged.status_code, 200)
        self.assertIn(finding_id, acknowledged.get_json()["acknowledged"])

        missing_rescan = self.client.post(
            f"/api/sessions/{session_id}/rescan",
            data={},
            content_type="multipart/form-data",
        )
        self.assertEqual(missing_rescan.status_code, 400)
        corrected_configs = io.BytesIO()
        with zipfile.ZipFile(corrected_configs, "w") as zip_file:
            zip_file.writestr("Configs/Api/settings.json", '{"Endpoint":"ready"}')
        corrected_configs.seek(0)
        rescanned = self.client.post(
            f"/api/sessions/{session_id}/rescan",
            data={"archive": (corrected_configs, "Configs.zip")},
            content_type="multipart/form-data",
        )
        self.assertEqual(rescanned.status_code, 200)
        self.assertTrue(rescanned.get_json()["rescan_completed"])

        blocked_local_upload = self.client.post(
            f"/api/sessions/{session_id}/upload",
            json={"approver": "Reviewer", "commit_message": "Codedrop10"},
        )
        self.assertEqual(blocked_local_upload.status_code, 400)
        self.assertIn("commits are disabled", blocked_local_upload.get_json()["error"])

        audit = self.client.get(f"/api/sessions/{session_id}/audit")
        self.assertEqual(audit.status_code, 200)
        self.assertIn("session_created", [item["event"] for item in audit.get_json()])
        self.assertIn("baseline_changed", [item["event"] for item in audit.get_json()])
        self.assertEqual(
            self.client.get("/api/sessions/not-a-session").status_code,
            404,
        )

    def test_release_without_review_areas_requires_explicit_skips(self):
        working_copy = Path(self.temp.name) / "SVN checkout"
        (working_copy / ".svn").mkdir(parents=True)
        previous_binary = working_copy / "Codedrop9" / "Binaries" / "Api" / "service.exe"
        previous_binary.parent.mkdir(parents=True)
        previous_binary.write_bytes(b"previous")
        previous_config = working_copy / "Codedrop9" / "Configs" / "Api" / "settings.json"
        previous_config.parent.mkdir(parents=True)
        previous_config.write_text('{"port": 80}', encoding="utf-8")

        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": str(working_copy),
        }).get_json()
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as zip_file:
            zip_file.writestr("Codedrop10/README.txt", "No review folders in this release.")
        archive.seek(0)
        started = self.client.post(
            f"/api/sessions/{created['id']}/source",
            data={"archive": (archive, "Codedrop10.zip")},
            content_type="multipart/form-data",
        )
        self.assertEqual(started.status_code, 202)
        for _ in range(100):
            response = self.client.get(f"/api/sessions/{created['id']}").get_json()
            if response["status"] != "analyzing":
                break
            time.sleep(0.02)

        self.assertEqual(response["status"], "ready", response["state"].get("error"))
        state = response["state"]
        self.assertEqual(state["areas_missing"], ["Binaries", "Configs"])
        self.assertEqual(state["config_scan_status"], "not run: Configs directory missing")
        self.assertFalse(state["rescan_completed"])
        self.assertEqual(
            {item["status"] for item in state["binaries"]},
            {"Removed"},
        )
        self.assertEqual(
            {item["status"] for item in state["configs"]},
            {"Removed"},
        )
        self.assertEqual(
            self.client.post(
                f"/api/sessions/{created['id']}/review",
                json={"reviewed": True},
            ).status_code,
            400,
        )
        skipped_binary = self.client.post(
            f"/api/sessions/{created['id']}/skip-area",
            json={"area": "Binaries"},
        )
        self.assertEqual(skipped_binary.status_code, 200)
        skipped_config = self.client.post(
            f"/api/sessions/{created['id']}/skip-area",
            json={"area": "Configs"},
        )
        self.assertEqual(skipped_config.status_code, 200)
        self.assertEqual(
            set(skipped_config.get_json()["skipped_areas"]),
            {"Binaries", "Configs"},
        )
        self.assertEqual(
            self.client.post(
                f"/api/sessions/{created['id']}/skip-area",
                json={"area": "Configs"},
            ).status_code,
            200,
        )
        bundle_response = self.client.get(
            f"/api/sessions/{created['id']}/reports/Release-Review.zip"
        )
        self.assertEqual(bundle_response.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(bundle_response.data)) as bundle:
            notes = bundle.read("Codedrop10-Review-Notes.txt").decode("utf-8")
            self.assertIn("Binary comparison: not included; skip confirmed", notes)
            self.assertIn("Config comparison: not included; skip confirmed", notes)
            self.assertIn("All missing review-area skips have been confirmed.", notes)
        bundle_response.close()

    def test_source_endpoint_rejects_svn_source_paths(self):
        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": "https://svn.example/repo/CodeDrops",
        }).get_json()
        with patch("app.svn_export") as export:
            response = self.client.post(
                f"/api/sessions/{created['id']}/source",
                data={
                    "source_mode": "svn",
                    "source_svn_url": "https://svn.example/repo/CodeDrops/Codedrop09",
                },
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Only ZIP uploads", response.get_json()["error"])
        export.assert_not_called()

    def test_local_folder_comparison_generates_reports_without_svn_commands(self):
        root = Path(self.temp.name)
        working_copy = root / "Release history"
        working_copy.mkdir()
        for drop in ("Codedrop9", "Codedrop8", "Codedrop6", "Codedrop5"):
            binaries = working_copy / drop / "Binaries" / "Api"
            binaries.mkdir(parents=True)
            (binaries / "service.exe").write_bytes(drop.encode())
        baseline_configs = working_copy / "Codedrop9" / "Configs" / "Api"
        baseline_configs.mkdir(parents=True)
        (baseline_configs / "settings.json").write_text('{"port": 9}', encoding="utf-8")
        legacy = working_copy / "Codedrop9" / "Binaries" / "Legacy"
        legacy.mkdir(parents=True)
        (legacy / "retired.exe").write_bytes(b"legacy component binary")

        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": str(working_copy),
        })
        self.assertEqual(created.status_code, 201, created.get_json())
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as zip_file:
            zip_file.writestr("Codedrop10/Binaries/Api/service.exe", b"Codedrop10")
        archive.seek(0)

        with patch("app.svn_directory_entries", side_effect=AssertionError("SVN CLI must not run")), \
             patch("app.svn_component_revision", side_effect=AssertionError("SVN CLI must not run")), \
             patch("app.svn_export", side_effect=AssertionError("SVN CLI must not run")):
            started = self.client.post(
                f"/api/sessions/{created.get_json()['id']}/source",
                data={"archive": (archive, "Codedrop10.zip")},
                content_type="multipart/form-data",
            )
            self.assertEqual(started.status_code, 202)
            for _ in range(100):
                response = self.client.get(
                    f"/api/sessions/{created.get_json()['id']}"
                ).get_json()
                if response["status"] != "analyzing":
                    break
                time.sleep(0.02)

        state = response["state"]
        self.assertEqual(response["status"], "ready", state.get("error"))
        self.assertEqual(state["target_type"], "local_folder")
        self.assertEqual(state["baselines"]["Binaries/Api"]["code_drop"], "Codedrop9")
        self.assertEqual(state["baselines"]["Binaries/Legacy"]["code_drop"], "Codedrop9")
        removed_component_files = [
            item for item in state["binaries"]
            if item["component"] == "Legacy" and item["status"] == "Removed"
        ]
        self.assertEqual([item["file"] for item in removed_component_files], ["retired.exe"])
        self.assertEqual(state["config_scan_status"], "not run: Configs directory missing")
        self.assertEqual(
            [(item["file"], item["status"]) for item in state["configs"]],
            [("Api/settings.json", "Removed")],
        )
        self.assertEqual(
            state["history_releases"],
            ["Codedrop9", "Codedrop8", "Codedrop6", "Codedrop5", "Codedrop10"],
        )
        self.assertFalse(state["rescan_completed"])
        self.assertNotIn("Configs", state["skipped_areas"])
        skipped_configs = self.client.post(
            f"/api/sessions/{created.get_json()['id']}/skip-area",
            json={"area": "Configs"},
        )
        self.assertEqual(skipped_configs.status_code, 200)
        self.assertIn("Configs", skipped_configs.get_json()["skipped_areas"])
        self.assertIn("Codedrop5", state["release_inventory"])
        self.assertIn("Codedrop9", state["release_inventory"])
        unavailable_drop = next(
            item for item in state["release_coverage"]
            if item["code_drop"] == "Codedrop7"
            and item["component_path"] == "Binaries/Api"
        )
        self.assertEqual(unavailable_drop["comparison"], "Release unavailable")
        for report in ("Binaries.xlsx", "Configs.xlsx"):
            result = self.client.get(
                f"/api/sessions/{created.get_json()['id']}/reports/{report}"
            )
            self.assertEqual(result.status_code, 200)
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(result.data), read_only=True)
            self.assertIn("Release Coverage", workbook.sheetnames)
            coverage = list(workbook["Release Coverage"].values)
            self.assertTrue(any(
                row[0] == "Codedrop7" and row[2] == "Release unavailable"
                for row in coverage[1:]
            ))
            if report == "Configs.xlsx":
                summary = {row[0].value: row[1].value for row in workbook["Scan Summary"].iter_rows(min_row=2)}
                self.assertEqual(summary["Configs directory present"], "No")
                self.assertEqual(summary["Dummy scan status"], "not run: Configs directory missing")
                self.assertNotIn("Removed files", summary)
            else:
                comparison = list(workbook["Binary Comparison"].values)
                self.assertFalse(any(row[-1] == "Removed" for row in comparison[1:]))
            workbook.close()
            result.close()
        blocked = self.client.post(
            f"/api/sessions/{created.get_json()['id']}/upload",
            json={"commit_message": "Codedrop10", "approver": "Test"},
        )
        self.assertEqual(blocked.status_code, 400)
        self.assertIn("commits are disabled", blocked.get_json()["error"])

    def test_local_working_copy_archived_code_drop_is_extracted_and_compared_without_mutation(self):
        root = Path(self.temp.name)
        working_copy = root / "SVN checkout"
        (working_copy / ".svn").mkdir(parents=True)
        archive_path = working_copy / "Codedrop9.zip"
        with zipfile.ZipFile(archive_path, "w") as zip_file:
            zip_file.writestr("Codedrop9/Binaries/Api/service.exe", b"release9")
            zip_file.writestr("Codedrop9/Binaries/Api/retired.dll", b"retired")
            zip_file.writestr("Codedrop9/Configs/Api/settings.json", '{"port": 9}')
        original_archive = archive_path.read_bytes()
        baseline8 = working_copy / "Codedrop8" / "Binaries" / "Api"
        baseline8.mkdir(parents=True)
        (baseline8 / "service.exe").write_bytes(b"Codedrop10")

        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": str(working_copy),
        }).get_json()
        upload = io.BytesIO()
        with zipfile.ZipFile(upload, "w") as zip_file:
            zip_file.writestr("Codedrop10/Binaries/Api/service.exe", b"Codedrop10")
        upload.seek(0)
        started = self.client.post(
            f"/api/sessions/{created['id']}/source",
            data={
                "archive": (upload, "Codedrop10.zip"),
                "svn_username": "test-user",
                "svn_password": "test-password",
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(started.status_code, 202)
        for _ in range(100):
            response = self.client.get(f"/api/sessions/{created['id']}").get_json()
            if response["status"] != "analyzing":
                break
            time.sleep(0.02)

        self.assertEqual(response["status"], "ready", response["state"].get("error"))
        state = response["state"]
        self.assertEqual(state["baselines"]["Binaries/Api"]["code_drop"], "Codedrop9")
        self.assertEqual(archive_path.read_bytes(), original_archive)
        binary_report = self.client.get(
            f"/api/sessions/{created['id']}/reports/Binaries.xlsx"
        )
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(binary_report.data), read_only=True)
        history = list(workbook["All Release Comparisons"].values)
        self.assertNotIn(
            ("Codedrop9", "Api", "retired.dll", "Missing from current"),
            [row[:4] for row in history],
        )
        self.assertIn(
            ("Codedrop8", "Api", "service.exe", "Unchanged"),
            [row[:4] for row in history],
        )
        workbook.close()
        binary_report.close()
        config_report = self.client.get(
            f"/api/sessions/{created['id']}/reports/Configs.xlsx"
        )
        workbook = load_workbook(io.BytesIO(config_report.data), read_only=True)
        config_history = list(workbook["All Release Comparisons"].values)
        self.assertNotIn(
            ("Codedrop9", "Api", "settings.json", "Missing from current"),
            [row[:4] for row in config_history],
        )
        workbook.close()
        config_report.close()

    def test_nested_release_archives_are_discovered_and_compared_in_release_order(self):
        root = Path(self.temp.name)
        working_copy = root / "SVN checkout"
        (working_copy / ".svn").mkdir(parents=True)
        release9 = working_copy / "Codedrop9" / "Binaries" / "Api"
        release9.mkdir(parents=True)
        (release9 / "service.exe").write_bytes(b"release9")
        same_release_archive = working_copy / "Codedrop9" / "packages" / "Codedrop9.zip"
        same_release_archive.parent.mkdir(parents=True)
        with zipfile.ZipFile(same_release_archive, "w") as zip_file:
            zip_file.writestr("Codedrop9/Configs/Api/settings.json", '{"port": 9}')
        nested_archive = working_copy / "Codedrop9" / "packages" / "historic.zip"
        with zipfile.ZipFile(nested_archive, "w") as zip_file:
            zip_file.writestr("Codedrop8/Binaries/Api/service.exe", b"release8")
            zip_file.writestr("Codedrop8/Configs/Api/settings.json", '{"port": 8}')
        nested_archive7 = working_copy / "Codedrop9" / "packages" / "Codedrop7.zip"
        with zipfile.ZipFile(nested_archive7, "w") as zip_file:
            zip_file.writestr("Binaries/Api/service.exe", b"release7")

        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": str(working_copy),
        }).get_json()
        upload = io.BytesIO()
        with zipfile.ZipFile(upload, "w") as zip_file:
            zip_file.writestr("Codedrop10/Binaries/Api/service.exe", b"release10")
        upload.seek(0)
        started = self.client.post(
            f"/api/sessions/{created['id']}/source",
            data={"archive": (upload, "Codedrop10.zip")},
            content_type="multipart/form-data",
        )
        self.assertEqual(started.status_code, 202)
        for _ in range(100):
            response = self.client.get(f"/api/sessions/{created['id']}").get_json()
            if response["status"] != "analyzing":
                break
            time.sleep(0.02)
        self.assertEqual(response["status"], "ready", response["state"].get("error"))
        state = response["state"]
        self.assertEqual(state["historical_code_drops"], [
            "Codedrop9", "Codedrop8", "Codedrop7",
        ])
        self.assertEqual(state["baselines"]["Binaries/Api"]["code_drop"], "Codedrop9")
        from openpyxl import load_workbook
        binary_report = self.client.get(
            f"/api/sessions/{created['id']}/reports/Binaries.xlsx"
        )
        workbook = load_workbook(io.BytesIO(binary_report.data), read_only=True)
        binary_rows = list(workbook["All Release Comparisons"].values)
        self.assertTrue(any(
            row[0] == "Codedrop8" and row[2] == "service.exe"
            for row in binary_rows[1:]
        ))
        workbook.close()
        binary_report.close()
        config_report = self.client.get(
            f"/api/sessions/{created['id']}/reports/Configs.xlsx"
        )
        workbook = load_workbook(io.BytesIO(config_report.data), read_only=True)
        config_rows = list(workbook["All Release Comparisons"].values)
        self.assertFalse(any(
            row[0] == "Codedrop8" and row[2] == "settings.json"
            for row in config_rows[1:]
        ))
        self.assertFalse(any(
            row[0] == "Codedrop9" and row[2] == "settings.json"
            for row in config_rows[1:]
        ))
        workbook.close()
        config_report.close()

    def test_svn_archived_code_drop_is_exported_then_compared(self):
        baseline_archive = io.BytesIO()
        with zipfile.ZipFile(baseline_archive, "w") as zip_file:
            zip_file.writestr("Codedrop9/Binaries/Api/service.exe", b"old")
        baseline_bytes = baseline_archive.getvalue()
        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": "https://svn.example/repo/CodeDrops",
        }).get_json()
        upload = io.BytesIO()
        with zipfile.ZipFile(upload, "w") as zip_file:
            zip_file.writestr("Codedrop10/Binaries/Api/service.exe", b"new")
        upload.seek(0)

        def export_archive(_url, destination, *_args):
            Path(destination).parent.mkdir(parents=True, exist_ok=True)
            Path(destination).write_bytes(baseline_bytes)

        with patch("app.svn_directory_entries", return_value=[{
            "name": "Codedrop9.zip", "kind": "file", "revision": "123",
        }]), patch("app.svn_export", side_effect=export_archive):
            started = self.client.post(
                f"/api/sessions/{created['id']}/source",
                data={
                    "archive": (upload, "Codedrop10.zip"),
                    "svn_username": "test-user",
                    "svn_password": "test-password",
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(started.status_code, 202)
            for _ in range(100):
                response = self.client.get(
                    f"/api/sessions/{created['id']}"
                ).get_json()
                if response["status"] != "analyzing":
                    break
                time.sleep(0.02)

        self.assertEqual(response["status"], "ready", response["state"].get("error"))
        state = response["state"]
        self.assertEqual(state["baselines"]["Binaries/Api"]["code_drop"], "Codedrop9")
        self.assertEqual(state["binaries"][0]["status"], "Changed")
        self.assertIn("Codedrop9", state["historical_code_drops"])

    def test_invalid_json_still_produces_a_configs_validation_report(self):
        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": "https://svn.example/repo/CodeDrops",
        }).get_json()
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as zip_file:
            zip_file.writestr("Codedrop10/Configs/Api/settings.json", '{"broken": }')
        archive.seek(0)
        with patch("app.svn_directory_entries", return_value=[]):
            started = self.client.post(
                f"/api/sessions/{created['id']}/source",
                data={
                    "archive": (archive, "Codedrop10.zip"),
                    "svn_username": "test-user",
                    "svn_password": "test-password",
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(started.status_code, 202)
            for _ in range(100):
                response = self.client.get(f"/api/sessions/{created['id']}").get_json()
                if response["status"] != "analyzing":
                    break
                time.sleep(0.02)
        self.assertEqual(response["status"], "validation_failed")
        state = response["state"]
        self.assertEqual(state["config_scan_status"], "blocked: invalid JSON")
        self.assertEqual(
            [(item["file"], item["status"]) for item in state["configs"]],
            [("Api/settings.json", "Invalid JSON")],
        )
        report = self.client.get(
            f"/api/sessions/{created['id']}/reports/Configs.xlsx"
        )
        self.assertEqual(report.status_code, 200)
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(report.data), read_only=True)
        validation = list(workbook["Validation Issues"].values)
        self.assertEqual(validation[1][0], "Invalid JSON")
        self.assertEqual(validation[1][1], "Codedrop10/Configs/Api/settings.json")
        summary_rows = list(workbook["Scan Summary"].values)
        summary = {row[0]: row[1] for row in summary_rows[1:]}
        self.assertEqual(summary["Dummy scan status"], "blocked: invalid JSON")
        self.assertEqual(summary["Screenshot captured"], "Yes (embedded in current config sheet image)")
        self.assertEqual(summary["Invalid JSON files"], 1)
        workbook.close()
        report.close()

    def test_config_screen_has_capture_screenshot_button(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Capture screenshot", response.get_data(as_text=True))

    def test_zip_upload_runs_review_and_persists_report(self):
        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": "https://svn.example/repo/CodeDrops",
        }).get_json()
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as zip_file:
            zip_file.writestr(
                "Codedrop10/Configs/Api/settings.json",
                '{"Database": {"Password": "password123"}}',
            )
            zip_file.writestr("Codedrop10/Binaries/Api/service.exe", b"demo binary")
        archive.seek(0)
        with patch("app.svn_directory_entries", return_value=[]):
            response = self.client.post(
                f"/api/sessions/{created['id']}/source",
                data={
                    "source_mode": "zip",
                    "archive": (archive, "Codedrop10.zip"),
                    "svn_username": "test-user",
                    "svn_password": "test-password",
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(response.status_code, 202)
            for _ in range(100):
                current = self.client.get(f"/api/sessions/{created['id']}").get_json()
                if current["status"] != "analyzing":
                    break
                time.sleep(0.02)
        self.assertEqual(current["status"], "ready", current["state"].get("error"))
        self.assertEqual(current["state"]["validation"]["failed"], 0)
        self.assertEqual(current["state"]["findings"][0]["severity"], "Blocker")
        report = self.client.get(f"/api/sessions/{created['id']}/reports/Binaries.xlsx")
        self.assertEqual(report.status_code, 200)
        report.close()
        self.assertFalse(current["state"]["rescan_completed"])
        reviewed = self.client.post(
            f"/api/sessions/{created['id']}/review", json={"reviewed": True}
        )
        self.assertEqual(reviewed.status_code, 200)
        blocked = self.client.post(f"/api/sessions/{created['id']}/upload", json={
            "approver": "Test Approver",
            "commit_message": "Codedrop10 release",
        })
        self.assertEqual(blocked.status_code, 400)
        self.assertIn("fresh Configs ZIP", blocked.get_json()["error"])

        corrected = io.BytesIO()
        with zipfile.ZipFile(corrected, "w") as zip_file:
            zip_file.writestr("Configs/Api/settings.json", '{"enabled": true}')
        corrected.seek(0)
        rescanned = self.client.post(
            f"/api/sessions/{created['id']}/rescan",
            data={"archive": (corrected, "Configs.zip")},
            content_type="multipart/form-data",
        )
        self.assertEqual(rescanned.status_code, 200, rescanned.get_json())
        self.assertTrue(rescanned.get_json()["rescan_completed"])
        self.assertEqual(len(rescanned.get_json()["finding_changes"]["fixed"]), 1)
        self.assertEqual(rescanned.get_json()["finding_changes"]["fixed"][0]["current_value"], "–")
        with patch("app.svn_directory_entries", return_value=[]), patch(
            "app.svn_import", return_value=("42", "Committed revision 42.")
        ):
            uploaded = self.client.post(f"/api/sessions/{created['id']}/upload", json={
                "approver": "Test Approver",
                "commit_message": "Codedrop10 release",
            })
        self.assertEqual(uploaded.status_code, 200, uploaded.get_json())
        self.assertEqual(uploaded.get_json()["upload"]["revision"], "42")

    def test_dummy_findings_identify_current_and_baseline_release_files(self):
        working_copy = Path(self.temp.name) / "Release history"
        baseline_config = working_copy / "Codedrop9" / "Configs" / "Api" / "settings.json"
        baseline_config.parent.mkdir(parents=True)
        baseline_config.write_text('{"Credential":"safe"}', encoding="utf-8")
        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": str(working_copy),
        })
        self.assertEqual(created.status_code, 201)
        session_id = created.get_json()["id"]
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as zip_file:
            zip_file.writestr(
                "Codedrop10/Configs/Api/settings.json",
                '{"Credential":"password123"}',
            )
        archive.seek(0)

        started = self.client.post(
            f"/api/sessions/{session_id}/source",
            data={"archive": (archive, "Codedrop10.zip")},
            content_type="multipart/form-data",
        )
        self.assertEqual(started.status_code, 202)
        for _ in range(100):
            response = self.client.get(f"/api/sessions/{session_id}").get_json()
            if response["status"] != "analyzing":
                break
            time.sleep(0.02)

        self.assertEqual(response["status"], "ready", response["state"].get("error"))
        finding = response["state"]["findings"][0]
        self.assertEqual(finding["severity"], "Blocker")
        self.assertEqual(finding["release"], "Codedrop10")
        self.assertEqual(finding["current_file"], "Codedrop10/Configs/Api/settings.json")
        self.assertEqual(finding["baseline_code_drop"], "Codedrop9")
        self.assertEqual(finding["baseline_file"], "Codedrop9/Configs/Api/settings.json")
        self.assertTrue(finding["baseline_file_exists"])

        report_response = self.client.get(
            f"/api/sessions/{session_id}/reports/Configs.xlsx"
        )
        self.assertEqual(report_response.status_code, 200)
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(report_response.data), read_only=True)
        row = next(workbook["Dummy Values"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(row[:4], (
            "Codedrop10",
            "Codedrop10/Configs/Api/settings.json",
            "Codedrop9",
            "Codedrop9/Configs/Api/settings.json",
        ))
        workbook.close()
        report_response.close()

    def test_svn_baselines_are_resolved_separately_for_each_area_and_component(self):
        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": "https://svn.example/repo/CodeDrops",
        }).get_json()
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as zip_file:
            zip_file.writestr("Codedrop10/Binaries/TMACServer/server.exe", b"new-server")
            zip_file.writestr("Codedrop10/Binaries/TMACServer/server.dll", b"new-server-dll")
            zip_file.writestr("Codedrop10/Binaries/TMACProxy/proxy.exe", b"new-proxy")
            zip_file.writestr(
                "Codedrop10/Configs/TMACServer/app/settings.json",
                '{"port": 9000}',
            )
        archive.seek(0)

        def component_revision(area_url, component, *_credentials):
            drop, area = area_url.rstrip("/").split("/")[-2:]
            if area == "Binaries" and component == "TMACServer" and drop == "Codedrop09":
                return None
            return {
                "Codedrop09": "90",
                "Codedrop08": "80",
                "Codedrop07": "70",
                "Codedrop06": "60",
            }.get(drop)

        def export_baseline(url, destination, *_options):
            destination.mkdir(parents=True, exist_ok=True)
            if "/Binaries/" in url:
                component = url.rstrip("/").split("/")[-1]
                name = "server.exe" if component == "TMACServer" else "proxy.exe"
                (destination / name).write_bytes(b"baseline")
                if component == "TMACServer":
                    (destination / "server.dll").write_bytes(b"baseline-dll")
            else:
                (destination / "app").mkdir()
                (destination / "app" / "settings.json").write_text('{"port": 8000}', encoding="utf-8")

        with patch("app.svn_directory_entries", return_value=[
            {"name": "Codedrop09", "kind": "dir", "revision": "99"},
            {"name": "Codedrop08", "kind": "dir", "revision": "88"},
            {"name": "Codedrop07", "kind": "dir", "revision": "77"},
            {"name": "Codedrop06", "kind": "dir", "revision": "66"},
        ]), patch("app.svn_component_revision", side_effect=component_revision), patch(
            "app.svn_export", side_effect=export_baseline
        ):
            started = self.client.post(
                f"/api/sessions/{created['id']}/source",
                data={
                    "source_mode": "zip",
                    "archive": (archive, "release.zip"),
                    "svn_username": "test-user",
                    "svn_password": "test-password",
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(started.status_code, 202)
            for _ in range(100):
                response = self.client.get(f"/api/sessions/{created['id']}").get_json()
                if response["status"] != "analyzing":
                    break
                time.sleep(0.02)
        self.assertEqual(response["status"], "ready", response["state"].get("error"))
        baselines = response["state"]["baselines"]
        self.assertEqual(baselines["Binaries/TMACServer"]["code_drop"], "Codedrop08")
        self.assertEqual(baselines["Configs/TMACServer"]["code_drop"], "Codedrop09")
        self.assertEqual(baselines["Binaries/TMACProxy"]["code_drop"], "Codedrop09")
        self.assertEqual(
            response["state"]["history_releases"],
            ["Codedrop09", "Codedrop08", "Codedrop07", "Codedrop06", "Codedrop10"],
        )
        server_history = next(
            item for item in response["state"]["version_history"]
            if item["component"] == "TMACServer"
        )
        self.assertEqual(server_history["versions"]["Codedrop08"], "N/A")
        self.assertTrue(any(
            item["component"] == "TMACServer" and item["file"] == "server.dll"
            for item in response["state"]["version_history"]
        ))
        self.assertNotIn("source_root", response["state"])
        self.assertNotIn("local_path", baselines["Configs/TMACServer"])

    def test_missing_svn_component_export_skips_to_an_older_baseline(self):
        created = self.client.post("/api/sessions", json={
            "folder_name": "Codedrop10",
            "target_url": "https://svn.example/repo/CodeDrops",
        }).get_json()
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as zip_file:
            zip_file.writestr(
                "Codedrop10/Binaries/TMACSDKProxy/service.dll",
                b"current binary",
            )
        archive.seek(0)

        def export_baseline(url, destination, *_options):
            if "/Codedrop5/Binaries/TMACSDKProxy" in url:
                raise WorkflowError(
                    "SVN command failed: svn: E170000: URL "
                    "'https://svn.example/repo/CodeDrops/Codedrop5/Binaries/"
                    "TMACSDKProxy' doesn't exist"
                )
            destination.mkdir(parents=True, exist_ok=True)
            (destination / "service.dll").write_bytes(b"older binary")

        with patch("app.svn_directory_entries", return_value=[
            {"name": "Codedrop5", "kind": "dir", "revision": "55"},
            {"name": "Codedrop4", "kind": "dir", "revision": "44"},
        ]), patch(
            "app.svn_component_revision",
            side_effect=lambda area_url, _component, *_credentials:
                "55" if "Codedrop5" in area_url else "44",
        ), patch("app.svn_export", side_effect=export_baseline):
            started = self.client.post(
                f"/api/sessions/{created['id']}/source",
                data={
                    "source_mode": "zip",
                    "archive": (archive, "release.zip"),
                    "svn_username": "test-user",
                    "svn_password": "test-password",
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(started.status_code, 202)
            for _ in range(100):
                response = self.client.get(
                    f"/api/sessions/{created['id']}"
                ).get_json()
                if response["status"] != "analyzing":
                    break
                time.sleep(0.02)

        self.assertEqual(response["status"], "ready", response["state"].get("error"))
        self.assertEqual(
            response["state"]["baselines"]["Binaries/TMACSDKProxy"]["code_drop"],
            "Codedrop4",
        )
        self.assertTrue(any(
            item["code_drop"] == "Codedrop5"
            and item["comparison"] == "Component not present in release"
            for item in response["state"]["release_coverage"]
        ))


if __name__ == "__main__":
    unittest.main()
