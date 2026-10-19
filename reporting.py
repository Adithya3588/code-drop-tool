from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


INK = "173B32"
GREEN = "2D765D"
PALE_GREEN = "E8F2EB"
PALE_RED = "FCE9E6"
PALE_YELLOW = "FFF3D9"
MUTED = "66756E"


def _safe_cell_value(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip().startswith("="):
        return "'" + value
    return value


def _append_row(sheet, values: list[Any]) -> None:
    sheet.append([_safe_cell_value(value) for value in values])


def _style_sheet(sheet, widths: dict[int, int] | None = None) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if widths:
        for index, width in widths.items():
            sheet.column_dimensions[get_column_letter(index)].width = width


def _safe_sheet_name(value: str, existing: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", value)[:31] or "Config"
    name = base
    counter = 2
    while name in existing:
        suffix = f"_{counter}"
        name = base[:31 - len(suffix)] + suffix
        counter += 1
    existing.add(name)
    return name


def _add_historical_comparison_sheet(
    workbook: Workbook,
    title: str,
    rows: list[dict[str, Any]],
    area: str,
) -> None:
    sheet = workbook.create_sheet(title)
    _append_row(sheet, [
        "Compared release", "Component", "File", "Outcome",
        "Baseline SHA-256", "Current SHA-256", "JSON key changes",
    ])
    for row in rows:
        if row["area"] != area or row["status"] == "Missing from current":
            continue
        _append_row(sheet, [
            row["code_drop"], row["component"], row["file"], row["status"],
            row["baseline_sha256"], row["current_sha256"], row["json_key_changes"],
        ])
        outcome = sheet.cell(sheet.max_row, 4)
        if row["status"] == "Missing from current":
            outcome.fill = PatternFill("solid", fgColor=PALE_RED)
        elif row["status"] == "Unchanged":
            outcome.fill = PatternFill("solid", fgColor=PALE_GREEN)
        else:
            outcome.fill = PatternFill("solid", fgColor=PALE_YELLOW)
    _style_sheet(sheet, {1: 20, 2: 24, 3: 58, 4: 24, 5: 68, 6: 68, 7: 20})


def _add_release_coverage_sheet(
    workbook: Workbook,
    rows: list[dict[str, Any]],
    area: str,
) -> None:
    sheet = workbook.create_sheet("Release Coverage")
    _append_row(sheet, [
        "Expected earlier release", "Component", "Comparison", "Details",
    ])
    for row in rows:
        if row["area"] != area:
            continue
        _append_row(sheet, [
            row["code_drop"], row["component"], row["comparison"], row["details"],
        ])
        outcome = sheet.cell(sheet.max_row, 3)
        if row["comparison"] == "Release unavailable":
            outcome.fill = PatternFill("solid", fgColor=PALE_RED)
        elif row["comparison"] == "Component not present in release":
            outcome.fill = PatternFill("solid", fgColor=PALE_YELLOW)
        else:
            outcome.fill = PatternFill("solid", fgColor=PALE_GREEN)
    _style_sheet(sheet, {1: 24, 2: 32, 3: 38, 4: 76})


def _diff_png(config: dict[str, Any]) -> BytesIO:
    lines = [f"{config['file']}  |  {config['status']}"]
    if config.get("baseline"):
        lines.append(
            f"Baseline: {config['baseline']['code_drop']}/{config['baseline']['path']}"
        )
    else:
        lines.append("Baseline: none (new component)")
    parse_error = config.get("current_validation_error") or config.get("baseline_validation_error")
    if parse_error:
        lines.append(
            f"JSON parse issue: {parse_error.get('message')} "
            f"(line {parse_error.get('line')}, column {parse_error.get('column')})"
        )
        context = (parse_error.get("context") or "").strip()
        if context:
            lines.extend(context.splitlines()[:6])
        diff_lines = []
    else:
        diff_lines = config.get("text_diff") or [
            f"{change['change']}: {change['key_path']}  {change['previous']}  ->  {change['current']}"
            for change in config["changes"]
        ]
        if not diff_lines:
            lines.append("No semantic JSON changes.")
    lines.extend(diff_lines[:26])
    if len(diff_lines) > 26:
        lines.append("… remaining lines are available in the searchable table below")
    font = ImageFont.load_default()
    width = 1200
    rendered = []
    for line in lines:
        while len(line) > 130:
            rendered.append(line[:130])
            line = "  " + line[130:]
        rendered.append(line)
    height = 32 + len(rendered) * 22
    image = Image.new("RGB", (width, height), "#fbfaf7")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 34), fill=f"#{INK}")
    for index, line in enumerate(rendered):
        color = "#ffffff" if index == 0 else f"#{MUTED}"
        if index > 1 and line.startswith("+") and not line.startswith("+++"):
            color = "#20714c"
        elif index > 1 and line.startswith("-") and not line.startswith("---"):
            color = "#a13e37"
        draw.text((16, 11 + index * 22), line, fill=color, font=font)
    output = BytesIO()
    image.save(output, format="PNG")
    output.seek(0)
    return output


def build_binaries_report(
    destination: Path,
    folder_name: str,
    binary_rows: list[dict[str, Any]],
    history_releases: list[str] | None = None,
    version_history: list[dict[str, Any]] | None = None,
    historical_comparisons: list[dict[str, Any]] | None = None,
    release_coverage: list[dict[str, Any]] | None = None,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Binary Comparison"
    comparison = sheet
    _append_row(comparison, [
        "Component", "File", "Baseline folder path", "Baseline version",
        f"Current version ({folder_name})", "Status",
    ])
    for row in binary_rows:
        if row["status"] == "Removed":
            continue
        _append_row(comparison, [
            row["component"], row["file"], row["baseline_path"],
            row["baseline_version"], row["current_version"], row["status"],
        ])
        status_cell = comparison.cell(comparison.max_row, 6)
        if row["status"] in {"New", "Changed", "Moved", "Version downgrade", "Hash changed, same version"}:
            status_cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)
        else:
            status_cell.fill = PatternFill("solid", fgColor=PALE_GREEN)
    _style_sheet(comparison, {1: 24, 2: 42, 3: 72, 4: 20, 5: 24, 6: 34})

    history = workbook.create_sheet("Version History")
    releases = history_releases or [folder_name]
    _append_row(history, ["Component", "File"] + [
        f"{release} (current)" if release == folder_name else release
        for release in releases
    ])
    release_columns = {release: index + 3 for index, release in enumerate(releases)}
    for row in version_history or []:
        values = [row["component"], row["file"]]
        values.extend(row.get("versions", {}).get(release, "–") for release in releases)
        _append_row(history, values)
        current_cell = history.cell(history.max_row, release_columns[folder_name])
        current_cell.fill = PatternFill("solid", fgColor=PALE_GREEN)
        baseline_column = release_columns.get(row.get("baseline_drop", ""))
        if baseline_column:
            history.cell(history.max_row, baseline_column).fill = PatternFill(
                "solid", fgColor="D8EBDD"
            )
    history_widths = {1: 24, 2: 42}
    history_widths.update({column: 20 for column in release_columns.values()})
    _style_sheet(history, history_widths)
    for column in release_columns.values():
        history.cell(1, column).fill = PatternFill("solid", fgColor=GREEN)
    _add_historical_comparison_sheet(
        workbook, "All Release Comparisons", historical_comparisons or [], "Binaries"
    )
    _add_release_coverage_sheet(workbook, release_coverage or [], "Binaries")
    workbook.save(destination)


def build_configs_report(
    destination: Path,
    folder_name: str,
    configs: list[dict[str, Any]],
    findings: list[dict[str, Any]],
    rules_hash: str,
    validation: dict[str, Any] | None = None,
    configs_present: bool = True,
    scan_status: str = "completed",
    historical_comparisons: list[dict[str, Any]] | None = None,
    release_coverage: list[dict[str, Any]] | None = None,
) -> None:
    validation = validation or {
        "checked": 0,
        "valid": 0,
        "failed": 0,
        "errors": [],
        "duplicates": [],
    }
    visible_configs = [config for config in configs if config["status"] != "Removed"]
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _append_row(summary, ["Config file", "Baseline path", "Status", "Changed keys"])
    for config in visible_configs:
        baseline = config.get("baseline")
        _append_row(summary, [
            config["file"],
            f"{baseline['code_drop']}/{baseline['path']}" if baseline else "No baseline",
            config["status"], len(config["changes"]),
        ])
    _style_sheet(summary, {1: 54, 2: 76, 3: 18, 4: 16})

    scan_summary = workbook.create_sheet("Scan Summary")
    _append_row(scan_summary, ["Metric", "Result"])
    _append_row(scan_summary, ["Release", folder_name])
    _append_row(scan_summary, ["Configs directory present", "Yes" if configs_present else "No"])
    _append_row(scan_summary, ["Files compared", len(visible_configs)])
    _append_row(scan_summary, ["Added files", sum(item["status"] == "New" for item in visible_configs)])
    _append_row(scan_summary, ["Changed files", sum(item["status"] == "Changed" for item in visible_configs)])
    _append_row(scan_summary, ["Unchanged files", sum(item["status"] == "Unchanged" for item in visible_configs)])
    _append_row(scan_summary, [
        "Files with invalid baseline JSON",
        sum(bool(item.get("baseline_validation_error")) for item in visible_configs),
    ])
    _append_row(scan_summary, ["JSON files checked", validation.get("checked", 0)])
    _append_row(scan_summary, ["Valid JSON files", validation.get("valid", 0)])
    _append_row(scan_summary, ["Invalid JSON files", validation.get("failed", 0)])
    _append_row(scan_summary, ["Duplicate JSON keys", len(validation.get("duplicates", []))])
    _append_row(scan_summary, ["Dummy scan status", scan_status])
    _append_row(scan_summary, ["Screenshot captured", "Yes (embedded in current config sheet image)"])
    _append_row(scan_summary, ["Dummy findings", len(findings)])
    _append_row(scan_summary, ["Blockers", sum(item["severity"] == "Blocker" for item in findings)])
    _append_row(scan_summary, ["Warnings", sum(item["severity"] == "Warning" for item in findings)])
    _append_row(scan_summary, ["Rules SHA-256", rules_hash or "Not run"])
    _style_sheet(scan_summary, {1: 32, 2: 80})

    validation_sheet = workbook.create_sheet("Validation Issues")
    _append_row(validation_sheet, ["Type", "File", "Line", "Column", "Key path", "Message", "Context"])
    for error in validation.get("errors", []):
        _append_row(validation_sheet, [
            "Invalid JSON", f"{folder_name}/Configs/{error['file']}",
            error["line"], error["column"], "",
            error["message"], error["context"],
        ])
    for duplicate in validation.get("duplicates", []):
        _append_row(validation_sheet, [
            "Duplicate key", duplicate["file"], "", "", duplicate["key_path"],
            "Duplicate JSON object key; review the last-value-wins parse result.", "",
        ])
    for config in configs:
        issue = config.get("baseline_validation_error")
        baseline = config.get("baseline")
        if issue and baseline:
            _append_row(validation_sheet, [
                "Invalid baseline JSON",
                f"{baseline['code_drop']}/Configs/{config['file']}",
                issue["line"], issue["column"], "",
                issue["message"], "Baseline contents were not included; file hashes are reported instead.",
            ])
    _style_sheet(validation_sheet, {1: 18, 2: 48, 3: 10, 4: 10, 5: 42, 6: 72, 7: 70})

    names: set[str] = {"Summary"}
    for config in visible_configs:
        sheet = workbook.create_sheet(
            _safe_sheet_name(config["file"].replace("/", "_").replace(".json", ""), names)
        )
        _append_row(sheet, ["Current path", f"{folder_name}/Configs/{config['file']}"])
        baseline = config.get("baseline")
        _append_row(sheet, [
            "Baseline path",
            f"{baseline['code_drop']}/{baseline['path']}/"
            f"{config['file'].split('/', 1)[-1]}" if baseline else "No baseline",
        ])
        sheet.append([])
        image_lines = max(len(config["changes"]), len(config.get("text_diff", [])))
        image_height = min(700, max(100, 22 * (min(image_lines, 26) + 3)))
        table_header = max(9, (image_height + 19) // 20 + 5)
        sheet.cell(table_header, 1, _safe_cell_value("Key path"))
        sheet.cell(table_header, 2, _safe_cell_value("Previous value"))
        sheet.cell(table_header, 3, _safe_cell_value("Current value"))
        sheet.cell(table_header, 4, _safe_cell_value("Change"))
        for cell in sheet[table_header]:
            cell.fill = PatternFill("solid", fgColor=INK)
            cell.font = Font(color="FFFFFF", bold=True)
        for target_row, change in enumerate(config["changes"], start=table_header + 1):
            for column, value in enumerate([
                change["key_path"], change["previous"], change["current"], change["change"],
            ], start=1):
                sheet.cell(target_row, column, _safe_cell_value(value))
            color = PALE_GREEN if change["change"] == "Added" else (
                PALE_RED if change["change"] == "Removed" else PALE_YELLOW
            )
            for cell in sheet[target_row]:
                cell.fill = PatternFill("solid", fgColor=color)
        image = SheetImage(_diff_png(config))
        image.width = 800
        image.height = image_height
        sheet.add_image(image, "A4")
        _style_sheet(sheet, {1: 32, 2: 60, 3: 60, 4: 18})
        sheet.freeze_panes = f"A{table_header + 1}"
        sheet.auto_filter.ref = (
            f"A{table_header}:D{max(table_header, table_header + len(config['changes']))}"
        )

    dummy = workbook.create_sheet("Dummy Values")
    _append_row(dummy, [
        "Issue in release", "Current file", "Baseline code drop", "Baseline file",
        "Key path", "Value", "Matched pattern", "Severity",
    ])
    for finding in findings:
        _append_row(dummy, [
            finding.get("release", folder_name),
            finding.get("current_file", f"{folder_name}/Configs/{finding['file']}"),
            finding.get("baseline_code_drop") or "No baseline",
            (
                finding.get("baseline_file")
                if finding.get("baseline_file_exists")
                else (
                    f"No matching file (expected {finding.get('baseline_file')})"
                    if finding.get("baseline_file")
                    else "No matching baseline file"
                )
                if finding.get("baseline_code_drop")
                else "No component baseline"
            ),
            finding["key_path"], finding["value"], finding["pattern"], finding["severity"],
        ])
        dummy.cell(dummy.max_row, 8).fill = PatternFill(
            "solid", fgColor=PALE_RED if finding["severity"] == "Blocker" else PALE_YELLOW
        )
    _append_row(dummy, [])
    _append_row(dummy, ["Rules SHA-256", rules_hash])
    _style_sheet(dummy, {
        1: 22, 2: 64, 3: 22, 4: 64, 5: 38, 6: 46, 7: 30, 8: 18,
    })
    _add_historical_comparison_sheet(
        workbook, "All Release Comparisons", historical_comparisons or [], "Configs"
    )
    _add_release_coverage_sheet(workbook, release_coverage or [], "Configs")
    workbook.save(destination)
